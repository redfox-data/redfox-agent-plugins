#!/usr/bin/env python3
"""
YouTube Transcript Extractor — paste a YouTube link, get the video transcript
=========================================
Endpoint: POST /story/api/youtube/transcript

Usage:
    python3 extract.py "https://www.youtube.com/watch?v=dQw4w9WgXcQ"
    python3 extract.py "dQw4w9WgXcQ" --language "zh,en"
    python3 extract.py "https://youtu.be/dQw4w9WgXcQ" --timestamp   # with timestamps
    python3 extract.py "URL" --excel                                 # also export Excel
    python3 extract.py "URL" --json --no-metadata
"""

import argparse
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path

try:
    import requests
    HAS_REQUESTS = True
except ImportError:
    HAS_REQUESTS = False

try:
    from deep_translator import GoogleTranslator
    HAS_TRANSLATOR = True
except ImportError:
    HAS_TRANSLATOR = False

# ─── Configuration ──────────────────────────────────────────────────────────────
API_URL = "https://redfox.hk/story/api/youtube/transcript"
CONFIG_FILE = Path.home() / ".redfox" / "apis" / "redfox.json"
ENV_KEY = "REDFOX_API_KEY"

DEFAULT_OUTPUT_DIR = Path.home() / "Downloads" / "RedfoxYoutubeDigest"
DEFAULT_LANGUAGE = "zh,en,asr"
MAX_RETRIES = 3

# Chinese language code prefixes (the API language field may be zh / zh-CN / zh-TW etc.)
ZH_PREFIX = ("zh",)

# ─── Redfox channel attribution ──
# Appends the running AI client to the API `source` field for per-platform stats.
# Priority: REDFOX_CHANNEL env var (documented in SKILL.md) > client env markers > "github".
SOURCE_BASE = "YouTube提文案"


def detect_channel():
    """Best-effort detection of the AI client that is running this skill."""
    ch = os.environ.get("REDFOX_CHANNEL", "").strip().lower()
    if ch:
        return ch
    markers = (
        ("claude", ("CLAUDECODE", "CLAUDE_CODE_ENTRYPOINT", "CLAUDE_PLUGIN_ROOT")),
        ("cursor", ("CURSOR_TRACE_ID", "CURSOR_AGENT", "CURSOR_PLUGIN_ROOT")),
        ("codex", ("CODEX_SANDBOX", "CODEX_PLUGIN_ROOT", "CODEX_HOME")),
        ("gemini", ("GEMINI_CLI", "GEMINI_PLUGIN_ROOT")),
    )
    for name, keys in markers:
        if any(os.environ.get(k) for k in keys):
            return name
    return "github"


def build_source():
    canonical = {"claude": "Claude", "codex": "Codex", "cursor": "Cursor", "gemini": "Gemini", "github": "GitHub"}
    ch = detect_channel()
    return f"{SOURCE_BASE}-{canonical.get(ch, ch.capitalize())}"


SOURCE = build_source()

# ─── Terminal colors ─────────────────────────────────────────────────────────────
GREEN = "\033[92m"
YELLOW = "\033[93m"
RED = "\033[91m"
CYAN = "\033[96m"
BOLD = "\033[1m"
DIM = "\033[2m"
RESET = "\033[0m"


def is_chinese(language):
    """Check whether the API-returned language is Chinese."""
    lang_lower = (language or "").lower().replace("_", "-")
    return any(lang_lower.startswith(p) for p in ZH_PREFIX)


def translate_lines(lines, batch_size=30):
    """
    Batch-translate non-Chinese transcript lines into Chinese.
    Every batch_size lines are merged into one translation request to reduce API calls.
    Returns the translated lines (one-to-one with the input).
    """
    if not HAS_TRANSLATOR:
        warn("deep-translator is not installed, cannot auto-translate. Run: pip3 install deep-translator")
        return lines

    translator = GoogleTranslator(source="auto", target="zh-CN")
    translated = []

    for i in range(0, len(lines), batch_size):
        batch = lines[i:i + batch_size]
        # Merge with a special separator, split after translation
        separator = " \n⟪SEP⟫\n "
        merged = separator.join(batch)
        try:
            result = translator.translate(merged)
            parts = result.split("⟪SEP⟫")
            # Clean up extra whitespace
            parts = [p.strip() for p in parts]
            # If the split count does not match, fall back to line-by-line translation
            if len(parts) != len(batch):
                for line in batch:
                    try:
                        translated.append(translator.translate(line).strip())
                    except Exception:
                        translated.append(line)
            else:
                translated.extend(parts)
        except Exception as e:
            warn(f"Batch translation failed, falling back to line-by-line: {e}")
            for line in batch:
                try:
                    translated.append(translator.translate(line).strip())
                except Exception:
                    translated.append(line)

    return translated


def info(msg):
    print(f"{GREEN}[✓]{RESET} {msg}")

def warn(msg):
    print(f"{YELLOW}[!]{RESET} {msg}")

def error(msg):
    print(f"{RED}[✗]{RESET} {msg}")

def step(msg):
    print(f"{CYAN}[→]{RESET} {msg}")


# ─── API key management (CLI > env var > config file) ───────────────────────────
def get_api_key(cli_key=None):
    if cli_key:
        return cli_key
    env_key = os.environ.get(ENV_KEY)
    if env_key:
        return env_key
    if CONFIG_FILE.exists():
        try:
            data = json.loads(CONFIG_FILE.read_text())
            key = data.get("api_key")
            if key:
                return key
        except (json.JSONDecodeError, OSError):
            pass
    return None


def print_key_guide():
    error("No RedFox API Key detected. Configure it in any of these ways:")
    print(f"  1. Environment variable (recommended): export {ENV_KEY}=ak_your_key")
    print(f"  2. CLI argument: --api-key ak_your_key")
    print(f"  3. Config file: echo '{{\"api_key\":\"ak_your_key\"}}' > ~/.redfox/apis/redfox.json")
    print(f"  Sign up: https://redfox.hk/settings/api-keys?source=github")


# ─── Utilities ───────────────────────────────────────────────────────────────────
def fmt_ts(seconds):
    """seconds → MM:SS or HH:MM:SS"""
    try:
        seconds = int(float(seconds))
    except (ValueError, TypeError):
        return "00:00"
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h:02d}:{m:02d}:{s:02d}"
    return f"{m:02d}:{s:02d}"


def video_url_of(video_id):
    return f"https://www.youtube.com/watch?v={video_id}" if video_id else ""


# ─── Transcript API (incremental-delay retries + rate-limit handling) ───────────
def fetch_transcript(session, api_key, video_url, language, include_timestamp, send_metadata):
    """
    Call the transcript endpoint.
    Returns the data dict (video_id / language / transcript [/ metadata]) on success, None on failure.
    Retries cover three failure kinds: network errors, bad status codes, empty data; auth failures are not retried.
    """
    payload = {
        "videoUrl": video_url,
        "format": "json",
        "includeTimestamp": include_timestamp,
        "sendMetadata": send_metadata,
        "language": language,
        "source": SOURCE,
    }
    headers = {"REDFOX_API_KEY": api_key, "Content-Type": "application/json"}

    for attempt in range(MAX_RETRIES):
        try:
            resp = session.post(API_URL, json=payload, headers=headers, timeout=30)
            result = resp.json()

            # Gateway wrapper: {code, data, msg}
            if isinstance(result, dict) and "code" in result:
                code = result.get("code")
                if code in (200, 2000):
                    data = result.get("data") if isinstance(result.get("data"), dict) else result
                elif code in (3106, 3107):
                    error("API Key is invalid or expired, please check your configuration (3106/3107)")
                    return None
                elif code == 3108:
                    warn("Rate limit hit, retrying in 5 seconds…")
                    time.sleep(5)
                    continue
                else:
                    warn(f"API returned an error: code={code} msg={result.get('msg', '')}")
                    data = None
            else:
                # Not wrapped — the transcript structure directly
                data = result

            if isinstance(data, dict):
                transcript = data.get("transcript")
                # text-format fallback: a plain text string also counts as success
                if transcript or data.get("text"):
                    return data
            warn("Empty response or no transcript found (the video may have no captions)")

        except requests.exceptions.RequestException as e:
            warn(f"Network error: {e}")
        except (json.JSONDecodeError, ValueError):
            warn("Failed to parse response (not JSON)")

        if attempt < MAX_RETRIES - 1:
            delay = 0.5 * (attempt + 1)  # 0.5s → 1.0s incremental
            time.sleep(delay)

    return None


# ─── Output formatting ───────────────────────────────────────────────────────────
def render_lines(data, with_timestamp=True):
    """Render the transcript into a list of lines."""
    lines = []
    transcript = data.get("transcript") or []
    for seg in transcript:
        text = (seg.get("text") or "").strip()
        if not text:
            continue
        if with_timestamp:
            lines.append(f"[{fmt_ts(seg.get('start', 0))}] {text}")
        else:
            lines.append(text)
    if not lines and data.get("text"):
        lines = [t.strip() for t in str(data["text"]).splitlines() if t.strip()]
    return lines


def total_duration_of(data):
    transcript = data.get("transcript") or []
    if not transcript:
        return 0
    last = transcript[-1]
    try:
        return float(last.get("start", 0)) + float(last.get("duration", 0))
    except (ValueError, TypeError):
        return 0


def build_markdown(data, lines, with_timestamp, video_url_input):
    video_id = data.get("video_id", "")
    language = data.get("language", "")
    metadata = data.get("metadata") or {}
    title = metadata.get("title") or data.get("title") or ""
    author = metadata.get("author") or metadata.get("channel") or data.get("author") or ""

    md = ["# YouTube Video Transcript", ""]
    md.append("| Item | Value |")
    md.append("|------|------|")
    if title:
        md.append(f"| Title | {title} |")
    if author:
        md.append(f"| Channel | {author} |")
    md.append(f"| Video ID | {video_id} |")
    md.append(f"| Video URL | {video_url_of(video_id) or video_url_input} |")
    md.append(f"| Language | {language} |")
    md.append(f"| Segments | {len(data.get('transcript') or [])} |")
    md.append(f"| Total duration | {fmt_ts(total_duration_of(data))} |")
    md.append(f"| Timestamps | {'with' if with_timestamp else 'without'} |")
    md.append(f"| Extracted at | {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} |")
    md.append("")
    md.append("---")
    md.append("")
    md.extend(lines)
    md.append("")
    return "\n".join(md)


# ─── Excel export ────────────────────────────────────────────────────────────────
def export_excel(data, lines, fpath):
    """Export xlsx: Title / Duration / Video URL / Transcript (matches the current timestamp mode)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        return False

    video_id = data.get("video_id", "")
    metadata = data.get("metadata") or {}
    title = metadata.get("title") or data.get("title") or video_id

    wb = Workbook()
    ws = wb.active
    ws.title = "Transcript"
    ws.append(["Title", "Duration", "Video URL", "Transcript"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append([title, fmt_ts(total_duration_of(data)), video_url_of(video_id), "\n".join(lines)])

    for col, width in {"A": 40, "B": 10, "C": 52, "D": 100}.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(fpath)
    return True


# ─── Main flow ───────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        description="YouTube Transcript Extractor — paste a YouTube link/ID to extract the video transcript"
    )
    parser.add_argument("video_url", help="YouTube video URL (full or short link) or plain video ID")
    parser.add_argument("--language", default=DEFAULT_LANGUAGE,
                        help=f"Caption language priority, comma separated (default {DEFAULT_LANGUAGE}: Chinese track first, falls back to English, asr = auto-generated captions; the API only picks tracks, it does not translate)")
    parser.add_argument("--timestamp", action="store_true",
                        help="Output transcript with [MM:SS] timestamps (no timestamps by default)")
    parser.add_argument("--no-metadata", action="store_true",
                        help="Skip fetching video metadata (title/channel fetched by default)")
    parser.add_argument("--excel", action="store_true",
                        help="Also export Excel (.xlsx, requires openpyxl)")
    parser.add_argument("--json", action="store_true",
                        help="Print the raw JSON response to the terminal")
    parser.add_argument("--no-save", action="store_true",
                        help="Do not save the Markdown file")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR),
                        help=f"Output directory (default {DEFAULT_OUTPUT_DIR})")
    parser.add_argument("--no-translate", dest="translate", action="store_false",
                        help="Disable auto translation (non-Chinese captions are translated to Chinese by default)")
    parser.set_defaults(translate=True)
    parser.add_argument("--api-key", default=None, help="RedFox API key")
    args = parser.parse_args()

    if not HAS_REQUESTS:
        error("The requests library is missing. Run: pip3 install requests")
        sys.exit(1)

    api_key = get_api_key(args.api_key)
    if not api_key:
        print_key_guide()
        sys.exit(1)

    with_timestamp = args.timestamp
    step(f"Extracting transcript: {args.video_url} (language priority: {args.language})")

    session = requests.Session()
    data = fetch_transcript(
        session, api_key, args.video_url,
        language=args.language,
        include_timestamp=True,  # always request timestamps for total-duration calculation; rendering layer decides
        send_metadata=not args.no_metadata,
    )
    if data is None:
        error("Extraction failed: transcript still unavailable after retries")
        error("Possible reasons: video has no captions / video does not exist or is restricted / API quota exhausted")
        sys.exit(1)

    if args.json:
        print(json.dumps(data, ensure_ascii=False, indent=2))

    # ─── Summary ───
    video_id = data.get("video_id", "")
    language = data.get("language", "")
    transcript = data.get("transcript") or []
    metadata = data.get("metadata") or {}
    title = metadata.get("title") or data.get("title") or ""
    author = metadata.get("author") or metadata.get("channel") or data.get("author") or ""

    info(f"Extraction succeeded: video {video_id}")
    if title:
        print(f"  Title: {title}")
    if author:
        print(f"  Channel: {author}")
    print(f"  Language: {language} | Segments: {len(transcript)} | Total duration: {fmt_ts(total_duration_of(data))}")
    print(f"  URL: {video_url_of(video_id) or args.video_url}")

    # ─── Auto translation (non-Chinese captions → Chinese) ───
    need_translate = args.translate and not is_chinese(language)
    if need_translate:
        step(f"Caption language is {language}, auto-translating to Chinese…")

    # ─── Full transcript ───
    lines = render_lines(data, with_timestamp)
    if not lines:
        warn("Transcript is empty, nothing to show")
    else:
        if need_translate:
            lines = translate_lines(lines)
            info(f"Translation complete ({len(lines)} segments)")
        print(f"\n{DIM}{'─' * 60}{RESET}")
        print(f"{BOLD}Full transcript ({len(lines)} segments):{RESET}\n")
        for line in lines:
            print(line)
        print(f"{DIM}{'─' * 60}{RESET}\n")

    # ─── Save files ───
    if not args.no_save:
        out_dir = Path(args.output_dir).expanduser()
        out_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = f"{video_id or 'transcript'}_{stamp}.md"
        fpath = out_dir / fname
        fpath.write_text(build_markdown(data, lines, with_timestamp, args.video_url), encoding="utf-8")
        info(f"Transcript saved: {fpath}")

    # ─── Excel export ───
    if args.excel:
        out_dir = Path(args.output_dir).expanduser()
        out_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        xpath = out_dir / f"{video_id or 'transcript'}_{stamp}.xlsx"
        if export_excel(data, lines, xpath):
            info(f"Excel saved: {xpath}")
        else:
            warn("openpyxl is not installed, cannot export Excel. Run: pip3 install openpyxl")


if __name__ == "__main__":
    main()
